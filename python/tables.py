"""Tables 1-4 as one APA-formatted Word document (derivatives/tables/tables.docx)."""

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from scipy import stats
from statsmodels.stats.multitest import multipletests

from config import DATA, TABLES

ID = "participant_id"
MINUS = "−"
EDUCATION = {
    1: "High school or equivalent or lower",
    2: "Some college or associate degree",
    3: "Bachelor's degree",
    4: "Graduate or professional degree",
}
INCOME = {
    1: "Less than $10,000",
    2: "$10,000 to $24,999",
    3: "$25,000 to $49,999",
    4: "$50,000 to $74,999",
    5: "$75,000 to $99,999",
    6: "$100,000 to $199,999",
    7: "$200,000 or more",
}
RACE = {
    3: "Non-Hispanic White",
    1: "Non-Hispanic Black",
    2: "Hispanic or Latino",
    4: "More than one option selected",
    0: "Non-Hispanic Asian",
}
LABOR = {0: "Spontaneous", 1: "Induced", 2: "No labor"}


def read(rel):
    return pd.read_csv(
        DATA / rel, sep="," if rel.endswith(".csv") else "\t", low_memory=False
    )


def signed(x, digits=2):
    return f"{x:+.{digits}f}".replace("-", MINUS)


def pval(p):
    return "<.001" if p < 0.001 else f"{p:.3f}".lstrip("0")


# Data ----------------------------------------------------------------------------


def modality_frames(sample):
    dwi = read("covariates/v02_dwi_primary_covariates.tsv").rename(
        columns={"sex_m": "child_sex"}
    )
    smri = read("analysis_datasets/v02_global_cortical_morphometry_model_data.tsv")
    fmri = read("analysis_datasets/v02_global_functional_hdp_data.tsv")
    keep = lambda d, flag: d[d[ID].isin(sample.loc[sample[flag], ID])].drop_duplicates(
        ID
    )
    return (
        keep(dwi, "in_diffusion"),
        keep(smri, "in_structural"),
        keep(fmri, "in_functional"),
    )


def characteristics(sample):
    dwi, smri, fmri = modality_frames(sample)
    core = [
        ID,
        "primary_hdp",
        "ga_delivery",
        "child_sex",
        "candidate_age_weeks",
        "chronic_hypertension",
    ]
    d = pd.concat([smri[core], fmri[core], dwi[core]]).drop_duplicates(ID)
    d["pma"] = d.ga_delivery + d.candidate_age_weeks
    d["preterm"] = (d.ga_delivery < 37).astype(float)
    person = read("covariates/hdp_missingness_person_level.csv").drop_duplicates(ID)
    complications = read("covariates/v02_pregnancy_complications.tsv")
    d = d.merge(
        person[[ID, "site", "mother_age", "education", "income", "race_eth"]],
        on=ID,
        how="left",
    ).merge(
        complications[
            [
                ID,
                "gestational_hypertension",
                "preeclampsia_eclampsia",
                "other_complication",
            ]
        ],
        on=ID,
        how="left",
    )
    pregnancy = (
        read("phenotype/pex_bm_healthv2_preg.csv")
        .drop_duplicates(ID)[
            [
                ID,
                "pex_bm_healthv2_preg__outcome_001",
                "pex_bm_healthv2_preg__outcome_002",
            ]
        ]
        .copy()
    )
    pregnancy["labor"] = pregnancy.pex_bm_healthv2_preg__outcome_001.where(
        pregnancy.pex_bm_healthv2_preg__outcome_001.isin(LABOR)
    )
    pregnancy["cesarean"] = pregnancy.pex_bm_healthv2_preg__outcome_002.where(
        pregnancy.pex_bm_healthv2_preg__outcome_002.isin([0, 1])
    )
    d = d.merge(pregnancy[[ID, "labor", "cesarean"]], on=ID, how="left")
    diffusion = read("outcomes/diffusion/v02_global_cerebral_wm_fa_md.tsv")
    diffusion["md_1e3"] = diffusion.global_median_md * 1000
    network = read("outcomes/functional/v02_functional_network_metrics.tsv")
    motion = read("covariates/v02_functional_qc.tsv")
    imaging = [
        (
            dwi,
            diffusion.merge(dwi[[ID, "raw_neighbor_corr"]], on=ID),
            ["md_1e3", "global_median_fa", "global_median_mkt", "raw_neighbor_corr"],
        ),
        (
            smri,
            smri,
            [
                "global_area_weighted_thickness",
                "global_area_weighted_absolute_sulc",
                "brainswipes_summary_qc",
            ],
        ),
        (
            fmri,
            network.merge(motion, on=ID),
            [
                "modularity",
                "mean_participation_coefficient",
                "mean_network_connectivity",
                "mean_fd_post_censoring",
                "retained_duration_minutes",
            ],
        ),
    ]
    for members, values, cols in imaging:
        values = values[values[ID].isin(members[ID])].drop_duplicates(ID)
        d = d.merge(values[[ID] + cols], on=ID, how="left")
    return d, len(dwi), len(smri), len(fmri)


def smd(a, b, binary):
    a, b = a.dropna(), b.dropna()
    if binary:
        sd = np.sqrt((a.mean() * (1 - a.mean()) + b.mean() * (1 - b.mean())) / 2)
    else:
        sd = np.sqrt((a.var() + b.var()) / 2)
    return (a.mean() - b.mean()) / sd


def table1(d, n_dwi, n_smri, n_fmri):
    hdp, ctl = d[d.primary_hdp == 1], d[d.primary_hdp == 0]

    def mean_row(col, label, digits=1):
        fmt = lambda s: f"{s.mean():.{digits}f} ({s.std():.{digits}f})"
        return [
            label,
            fmt(hdp[col].dropna()),
            fmt(ctl[col].dropna()),
            signed(smd(hdp[col], ctl[col], False)),
        ]

    def count_row(col, label, smd_cell=True):
        fmt = lambda s: f"{int(s.sum())} ({100 * s.mean():.1f})"
        return [
            label,
            fmt(hdp[col].dropna()),
            fmt(ctl[col].dropna()),
            signed(smd(hdp[col], ctl[col], True)) if smd_cell else "−",
        ]

    def category_rows(col, levels):
        rows = []
        for code, label in levels.items():
            d[f"_{code}"] = d[col].eq(code).astype(float).where(d[col].notna())
            hdp_, ctl_ = d[d.primary_hdp == 1], d[d.primary_hdp == 0]
            fmt = lambda s: f"{int(s.sum())} ({100 * s.mean():.1f})"
            rows.append(
                [
                    "   " + label,
                    fmt(hdp_[f"_{code}"].dropna()),
                    fmt(ctl_[f"_{code}"].dropna()),
                    signed(smd(hdp_[f"_{code}"], ctl_[f"_{code}"], True)),
                ]
            )
        return rows

    section = lambda label: [label, "", "", ""]
    sites = d.groupby("site").size()
    rows = (
        [
            section("Participant and pregnancy characteristics"),
            mean_row("mother_age", "Maternal age, years"),
            section("Maternal education, n (%)"),
        ]
        + category_rows("education", EDUCATION)
        + [section("Total household income, n (%)")]
        + category_rows("income", INCOME)
        + [section("Maternal race and ethnicity, n (%)")]
        + category_rows("race_eth", RACE)
        + [
            mean_row("ga_delivery", "Gestational age at birth, weeks"),
            count_row("preterm", "Preterm birth (< 37 weeks), n (%)"),
        ]
        + [section("Labor onset, n (%)")]
        + category_rows("labor", LABOR)
        + [
            count_row("cesarean", "Cesarean delivery, n (%)"),
            count_row("child_sex", "Infant male sex, n (%)"),
            mean_row("candidate_age_weeks", "Postnatal age at scan, weeks"),
            mean_row("pma", "Postmenstrual age at scan, weeks"),
            [
                "Recruitment sites contributing, n",
                str(hdp.site.nunique()),
                str(ctl.site.nunique()),
                "−",
            ],
            [
                "Dyads per site, median (range)",
                f"{sites.median():.0f} ({sites.min()} to {sites.max()})",
                "",
                "−",
            ],
            section("Hypertension characteristics"),
            count_row(
                "gestational_hypertension", "Gestational hypertension, n (%)", False
            ),
            count_row(
                "preeclampsia_eclampsia", "Pre-eclampsia or eclampsia, n (%)", False
            ),
            count_row(
                "chronic_hypertension", "General history of hypertension, n (%)", False
            ),
            count_row(
                "other_complication", "Any co-occurring pregnancy complication, n (%)"
            ),
            section("Imaging characteristics"),
            section(f"Diffusion MRI (n = {n_dwi})"),
            mean_row("md_1e3", "   Global mean diffusivity, ×10⁻³ mm²/s", 2),
            mean_row("global_median_fa", "   Global fractional anisotropy", 3),
            mean_row("global_median_mkt", "   Global mean kurtosis tensor", 3),
            mean_row("raw_neighbor_corr", "   Neighboring-DWI correlation", 3),
            section(f"Structural MRI (n = {n_smri})"),
            mean_row(
                "global_area_weighted_thickness", "   Global cortical thickness, mm", 3
            ),
            mean_row(
                "global_area_weighted_absolute_sulc",
                "   Global absolute sulcal depth, mm",
                3,
            ),
            mean_row("brainswipes_summary_qc", "   BrainSwipes quality score", 3),
            section(f"Functional MRI (n = {n_fmri})"),
            mean_row("modularity", "   Network modularity", 3),
            mean_row(
                "mean_participation_coefficient", "   Mean participation coefficient", 3
            ),
            mean_row("mean_network_connectivity", "   Mean network connectivity", 4),
            mean_row("mean_fd_post_censoring", "   Mean framewise displacement, mm", 3),
            mean_row(
                "retained_duration_minutes", "   Retained low-motion data, minutes", 1
            ),
        ]
    )
    header = [
        "Characteristic",
        f"HDP\n(n = {len(hdp)})",
        f"Normotensive\n(n = {len(ctl)})",
        "SMD",
    ]
    return pd.DataFrame(rows, columns=header)


def table2(d, sample):
    variables = [
        ("primary_hdp", "Hypertensive disorder of pregnancy"),
        ("ga_delivery", "Gestational age at birth"),
        ("pma", "Postmenstrual age at scan"),
        ("global_median_md", "Mean diffusivity"),
        ("global_median_fa", "Fractional anisotropy"),
        ("global_median_mkt", "Mean kurtosis tensor"),
        ("global_area_weighted_thickness", "Cortical thickness"),
        ("global_area_weighted_absolute_sulc", "Absolute sulcal depth"),
        ("modularity", "Functional modularity"),
        ("mean_participation_coefficient", "Participation coefficient"),
        ("other_complication", "Any other pregnancy complication"),
    ]
    harmonized = sample.drop(columns=[c for c in sample if c.startswith("in_")])
    data = harmonized.merge(
        d[[ID, "primary_hdp", "ga_delivery", "pma", "other_complication"]], on=ID
    )
    cols = [c for c, _ in variables]
    k = len(cols)
    r = np.full((k, k), np.nan)
    p = np.full((k, k), np.nan)
    n = np.zeros((k, k), int)
    for i in range(k):
        for j in range(i):
            pair = data[[cols[i], cols[j]]].dropna()
            n[i, j] = len(pair)
            r[i, j], p[i, j] = stats.pearsonr(pair[cols[i]], pair[cols[j]])
    lower = np.tril_indices(k, -1)
    q = np.full((k, k), np.nan)
    valid = np.isfinite(
        p[lower]
    )  # undefined correlations stay blank and outside the FDR family
    q[lower[0][valid], lower[1][valid]] = multipletests(
        p[lower][valid], method="fdr_bh"
    )[1]
    cell = lambda i, j: (
        "—"
        if i == j
        else (
            ""
            if j > i
            else f"{r[i, j]:.2f}".replace("0.", ".").replace("-", MINUS)
            + ("*" if q[i, j] < 0.05 else "")
        )
    )
    table = pd.DataFrame(
        [[cell(i, j) for j in range(k)] for i in range(k)],
        columns=[str(i + 1) for i in range(k)],
    )
    table.insert(
        0, "Variable", [f"{i + 1}. {label}" for i, (_, label) in enumerate(variables)]
    )
    return table, n[lower]


# Word formatting -------------------------------------------------------------------


def borders(table):
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single" if side in ("top", "bottom") else "none")
        b.set(qn("w:sz"), "8")
        el.append(b)
    table._tbl.tblPr.append(el)


def rule_below(row):
    for cell in row.cells:
        b = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        b.append(bottom)
        cell._tc.get_or_add_tcPr().append(b)


def add_row(table, widths, values, size, bold=False, italic=False, indent=0.0):
    cells = table.add_row().cells
    for i, (value, width) in enumerate(zip(values, widths)):
        cells[i].width = Inches(width)
        par = cells[i].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.left_indent = Inches(indent if i == 0 else 0)
        run = par.add_run(str(value))
        run.font.name, run.font.size, run.bold, run.italic = (
            "Arial",
            Pt(size),
            bold,
            italic,
        )
    return table.rows[-1]


def section_row(table, widths, label, size):
    row = add_row(table, widths, [label] + [""] * (len(widths) - 1), size, italic=True)
    row.cells[0].merge(row.cells[-1])


def title(doc, number, text):
    for line, bold, italic in [(f"Table {number}", True, False), (text, False, True)]:
        run = doc.add_paragraph().add_run(line)
        run.bold, run.italic, run.font.name, run.font.size = (
            bold,
            italic,
            "Arial",
            Pt(12),
        )


def note(doc, text):
    par = doc.add_paragraph()
    run = par.add_run("Note.")
    run.italic, run.font.name, run.font.size = True, "Arial", Pt(10)
    run = par.add_run(" " + text)
    run.font.name, run.font.size = "Arial", Pt(10)


def new_page(doc, landscape):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape != (sec.page_width > sec.page_height):
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = sec.page_height, sec.page_width


def write_simple(doc, frame, widths, size, sections=()):
    table = doc.add_table(rows=0, cols=len(widths))
    borders(table)
    rule_below(add_row(table, widths, frame.columns, size, bold=True))
    for _, row in frame.iterrows():
        label = str(row.iloc[0])
        r = add_row(
            table,
            widths,
            [label.strip()] + list(row.iloc[1:]),
            size,
            indent=0.18 if label.startswith("   ") else 0,
        )
        r.cells[0].paragraphs[0].runs[0].bold = label in sections
    rule_below(table.rows[-1])


def write_table3(doc, results):
    widths = [2.35, 1.45, 1.75, 0.65, 0.65, 1.75, 0.65, 0.65]
    table = doc.add_table(rows=0, cols=len(widths))
    borders(table)
    top = add_row(table, widths, [""] * len(widths), 9, bold=True)
    top.cells[2].merge(top.cells[4]).paragraphs[0].runs[0].text = "Primary model"
    top.cells[5].merge(top.cells[7]).paragraphs[0].runs[
        0
    ].text = "Accounting for gestational age at birth"
    stats_cols = ["β [95% CI]", "p", "q"]
    rule_below(
        add_row(
            table, widths, ["Measure", "n (HDP/control)"] + stats_cols * 2, 9, bold=True
        )
    )
    for modality in ["Diffusion MRI", "Structural MRI", "Functional MRI"]:
        section_row(table, widths, modality, 9)
        for outcome in results[results.modality == modality].outcome.unique():
            cells = []
            for model in ("primary", "ga_adjusted"):
                x = results[
                    (results.outcome == outcome) & (results.model_set == model)
                ].iloc[0]
                cells += [
                    f"{signed(x.beta)} [{signed(x.ci_lo)}, {signed(x.ci_hi)}]",
                    pval(x.p),
                    pval(x.q),
                ]
            add_row(
                table,
                widths,
                [outcome, f"{x.n:,} ({x.n_hdp:,}/{x.n_control:,})"] + cells,
                9,
                indent=0.15,
            )
    rule_below(table.rows[-1])


def write_table4(doc, sensitivity, loso):
    widths = [3.55, 1.15, 1.55, 0.65]
    table = doc.add_table(rows=0, cols=len(widths))
    borders(table)
    rule_below(
        add_row(
            table, widths, ["Analysis", "n (HDP)", "β [95% CI]", "p"], 10, bold=True
        )
    )
    loso_row = [
        f"Leaving out one scanner at a time ({len(loso)} scanners)",
        f"{loso.n.min():,}–{loso.n.max():,} ({loso.n_hdp.min()}–{loso.n_hdp.max()})",
        f"{signed(loso.beta.min())} to {signed(loso.beta.max())}",
        pval(loso.p.max()) if loso.p.max() < 0.001 else "≤" + pval(loso.p.max()),
    ]
    for group, rows in sensitivity.groupby("group", sort=False):
        section_row(table, widths, group, 10)
        for _, x in rows.iterrows():
            add_row(
                table,
                widths,
                [
                    x.analysis,
                    f"{x.n:,} ({x.n_hdp:,})",
                    f"{signed(x.beta)} [{signed(x.ci_lo)}, {signed(x.ci_hi)}]",
                    pval(x.p),
                ],
                10,
                indent=0.15,
            )
            if x.analysis == "No harmonization, scanner fixed effects":
                add_row(table, widths, loso_row, 10, indent=0.15)
    rule_below(table.rows[-1])


def main():
    sample = pd.read_csv(TABLES / "analytic_sample.csv")
    d, n_dwi, n_smri, n_fmri = characteristics(sample)
    t1 = table1(d, n_dwi, n_smri, n_fmri)
    t2, pair_n = table2(d, sample)
    results = pd.read_csv(TABLES / "global_models.csv")
    sensitivity = pd.read_csv(TABLES / "sensitivity.csv")
    loso = pd.read_csv(TABLES / "leave_one_scanner_out.csv")
    t1.to_csv(TABLES / "table1.csv", index=False)
    t2.to_csv(TABLES / "table2.csv", index=False)

    doc = Document()
    title(doc, 1, "Sample characteristics by hypertensive disorder of pregnancy status")
    write_simple(
        doc,
        t1,
        [3.30, 1.55, 1.75, 0.70],
        10,
        sections={
            "Participant and pregnancy characteristics",
            "Hypertension characteristics",
            "Imaging characteristics",
        },
    )
    note(
        doc,
        "Values are M (SD) or n (%). Percentages exclude missing data. The sample includes all dyads "
        "contributing at least one MRI modality, and imaging characteristics are reported within each "
        "modality's analytic sample. Postnatal age is time since birth; postmenstrual age is gestational "
        f"age at birth plus postnatal age. Education was missing for {d.education.isna().sum()} dyads and "
        f"household income for {d.income.isna().sum()}. Dashes indicate characteristics that define group "
        "membership. Values are unadjusted. HDP = hypertensive disorder of pregnancy; SMD = standardized "
        "mean difference.",
    )

    new_page(doc, landscape=True)
    title(doc, 2, "Correlations among core study variables")
    write_simple(doc, t2, [2.45] + [0.58] * (len(t2.columns) - 1), 9)
    note(
        doc,
        "Values are Pearson correlations. Coefficients involving the two binary variables are point-biserial "
        f"or phi correlations. Correlations were computed pairwise, so n ranges from {pair_n.min():,} to "
        f"{pair_n.max():,}. Postmenstrual age at scan is gestational age at birth plus postnatal age. Imaging "
        "measures were harmonized across scanners using CovBat-GAM. *q < .05 after Benjamini-Hochberg false "
        f"discovery rate correction across all {len(pair_n)} correlations.",
    )

    new_page(doc, landscape=True)
    title(
        doc,
        3,
        "Associations of hypertensive disorders of pregnancy with global infant brain measures",
    )
    write_table3(doc, results)
    note(
        doc,
        "β values are standardized differences between the HDP and normotensive groups; positive values "
        "indicate higher values in the HDP group. Imaging measures were harmonized across scanners (scanner "
        "× software version) with CovBat-GAM. Primary models adjusted for infant sex, image quality, and a "
        "smooth function of postmenstrual age at scan; models accounting for gestational age additionally "
        "included a smooth function of gestational age at birth. Confidence intervals and p values use "
        "heteroskedasticity-robust standard errors. q values are Benjamini-Hochberg false discovery "
        "rate-adjusted across the seven measures within each model. HDP = hypertensive disorder of pregnancy.",
    )

    new_page(doc, landscape=False)
    title(doc, 4, "Sensitivity analyses of the cortical thickness difference")
    write_table4(doc, sensitivity, loso)
    note(
        doc,
        "β values are standardized differences in global cortical thickness between the HDP and "
        "normotensive groups at matched postmenstrual age (positive = thicker in HDP), from the primary model "
        "with the listed modification. Each analysis was re-harmonized on its own sample. The "
        "leave-one-scanner-out row gives the range of sample sizes and estimates across models. Within-HDP "
        "rows compare HDP subgroups (n = HDP infants; second number = infants in the first-named subgroup). "
        "Birth weight for gestational age is a sex-specific z score computed within this sample from "
        "caregiver-reported birth weight; fetal growth restriction and neonatal intensive care admission were "
        "caregiver-reported. Total brain volume was summed from BIBSNet tissue segmentations. "
        "Substance use was any reported use on the Timeline Followback from the last menstrual period "
        "onward. Multiple gestation was identified from reported reasons for cesarean delivery. "
        "The combined-adjustment model was restricted to infants with complete "
        "data on all covariates. The effect-modification row gives the HDP × infant sex interaction "
        "coefficient. Confidence intervals and p values use heteroskedasticity-robust standard errors. HDP = "
        "hypertensive disorder of pregnancy; NICU = neonatal intensive care unit.",
    )
    doc.save(TABLES / "tables.docx")


if __name__ == "__main__":
    main()
