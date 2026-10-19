import numpy as np
import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn

from report_data import model_results, sensitivity_results, write_report_inputs


@pytest.fixture
def reports(modules):
    m = modules("python/tables.py")
    sample, core, values = write_report_inputs(m)
    return m, sample, core, values


@pytest.mark.parametrize("value,digits,expected", [(1.234, 2, "+1.23"), (-.125, 3, "−0.125"), (0, 1, "+0.0")])
def test_signed_number_format(modules, value, digits, expected):
    assert modules("python/tables.py").signed(value, digits) == expected


@pytest.mark.parametrize("value,expected", [(.0001, "<.001"), (.001, ".001"), (.42, ".420"), (1, "1.000")])
def test_pvalue_format(modules, value, expected):
    assert modules("python/tables.py").pval(value) == expected


def test_smd_uses_sample_variance_and_ignores_missing(modules):
    m = modules("python/tables.py")
    assert m.smd(pd.Series([1, 3, np.nan]), pd.Series([0, 2]), False) == pytest.approx(1 / np.sqrt(2))
    assert m.smd(pd.Series([1, 1, 0, 0]), pd.Series([1, 0, 0, 0, np.nan]), True) == pytest.approx(.25 / np.sqrt(.21875))


def test_characteristics_respect_modality_membership_and_units(reports):
    m, sample, core, values = reports
    d, *counts = m.characteristics(sample)
    assert counts == [31, 31, 31]
    assert len(d) == 32
    assert d.participant_id.is_unique
    d = d.set_index("participant_id")
    assert pd.isna(d.loc["sub-synthetic000", "md_1e3"])
    assert pd.isna(d.loc["sub-synthetic001", "global_area_weighted_thickness"])
    assert pd.isna(d.loc["sub-synthetic002", "modularity"])
    assert d.loc["sub-synthetic003", "md_1e3"] == pytest.approx(values.global_median_md[3] * 1000)
    np.testing.assert_allclose(d.loc[core.participant_id, "pma"], core.ga_delivery + core.candidate_age_weeks)


def test_modality_membership_excludes_duplicates_and_nonmembers(reports):
    m, sample, _, _ = reports
    relative = "covariates/v02_dwi_primary_covariates.tsv"
    data = m.read(relative)
    outsider = data.iloc[[1]].assign(participant_id="sub-not-in-sample")
    pd.concat([data, data.iloc[[1]], outsider]).to_csv(m.DATA / relative, sep="\t", index=False)
    dwi, _, _ = m.modality_frames(sample)
    assert len(dwi) == 31
    assert dwi.participant_id.is_unique
    assert "sub-not-in-sample" not in set(dwi.participant_id)
    assert "child_sex" in dwi


def test_table1_counts_and_missing_denominators(reports):
    m, sample, _, _ = reports
    d, *counts = m.characteristics(sample)
    result = m.table1(d, *counts).set_index("Characteristic")
    assert result.columns.tolist() == ["HDP\n(n = 16)", "Normotensive\n(n = 16)", "SMD"]
    assert result.loc["   High school or equivalent or lower", "Normotensive\n(n = 16)"] == "7 (46.7)"
    assert result.loc["Gestational hypertension, n (%)", "HDP\n(n = 16)"] == "16 (100.0)"
    assert "Diffusion MRI (n = 31)" in result.index


def test_table2_pairwise_missingness_direction_and_fdr(reports):
    m, sample, _, _ = reports
    d, *_ = m.characteristics(sample)
    # Hand-constructed perfect negative association; must be significant after 55 tests.
    sample["global_median_fa"] = -sample.global_median_md
    sample.loc[0, "global_median_mkt"] = np.nan
    table, counts = m.table2(d, sample)
    assert table.shape == (11, 12)
    assert table.iloc[4]["4"] == "−1.00*"
    assert table.iloc[4]["5"] == "—"
    assert table.iloc[4]["6"] == ""
    assert len(counts) == 55
    assert counts.min() == 31
    assert counts.max() == 32


def test_word_formatting_and_page_orientation(modules):
    m = modules("python/tables.py")
    doc = Document()
    m.title(doc, 2, "Synthetic table")
    m.note(doc, "Synthetic data only.")
    m.new_page(doc, True)
    assert doc.sections[-1].page_width > doc.sections[-1].page_height
    m.new_page(doc, True)
    m.new_page(doc, False)
    assert doc.sections[-1].page_width < doc.sections[-1].page_height
    assert doc.paragraphs[0].runs[0].bold
    assert doc.paragraphs[1].runs[0].italic
    assert doc.paragraphs[2].runs[0].italic
    frame = pd.DataFrame([["Section", ""], ["   Indented", "3"]], columns=["Measure", "Value"])
    m.write_simple(doc, frame, [3, 1], 10, sections={"Section"})
    table = doc.tables[0]
    assert table.cell(1, 0).paragraphs[0].runs[0].bold
    assert table.cell(2, 0).text == "Indented"
    assert table.cell(2, 0).paragraphs[0].paragraph_format.left_indent > 0
    assert table._tbl.tblPr.find(qn("w:tblBorders")) is not None
    assert table.cell(0, 0)._tc.tcPr.find(qn("w:tcBorders")) is not None


def test_model_tables_contain_estimates_and_scanner_ranges(modules):
    m = modules("python/tables.py")
    doc = Document()
    m.write_table3(doc, model_results())
    loso = pd.DataFrame({"n": [20, 24], "n_hdp": [8, 10], "beta": [.1, .3], "p": [.01, .03]})
    m.write_table4(doc, sensitivity_results(), loso)
    text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "+0.20 [+0.10, +0.30]" in text
    assert "32 (16/16)" in text
    assert "20–24 (8–10)" in text
    assert "≤.030" in text
    assert "Leaving out one scanner at a time (2 scanners)" in text


def test_tables_main_writes_reopenable_document_and_csvs(reports):
    m, _, _, _ = reports
    model_results().to_csv(m.TABLES / "global_models.csv", index=False)
    sensitivity_results().to_csv(m.TABLES / "sensitivity.csv", index=False)
    pd.DataFrame({"n": [20, 24], "n_hdp": [8, 10], "beta": [.1, .3], "p": [.0001, .0002]}).to_csv(
        m.TABLES / "leave_one_scanner_out.csv", index=False)
    m.main()
    doc = Document(m.TABLES / "tables.docx")
    assert len(doc.tables) == 4
    assert [p.text for p in doc.paragraphs if p.text.startswith("Table ")] == ["Table 1", "Table 2", "Table 3", "Table 4"]
    assert pd.read_csv(m.TABLES / "table1.csv").shape[1] == 4
    assert pd.read_csv(m.TABLES / "table2.csv").shape == (11, 12)
